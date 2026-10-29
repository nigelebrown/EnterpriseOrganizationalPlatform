"""HTTP-level acceptance checks through the actual API, governance and SQLite persistence.
SQL Server integration, OIDC issuer, AG failover and security penetration tests are separate release gates.
"""
from datetime import date
from uuid import uuid4
from jcf_platform import models as m

def loc(code,name='Division A',valid='2020-01-01',**kw):
    return dict(code=code,name=name,locationType='DIVISION',validFrom=valid,**kw)
def section(code,name):return dict(code=code,name=name,validFrom='2020-01-01')
def attach(locationId,definitionId,**kw):return dict(locationId=locationId,sectionDefinitionId=definitionId,validFrom='2020-01-01',**kw)
def getid(r):assert r.status_code==200,r.text;return r.json()['entityId']

def test_health_and_admin(client):
    assert client.get('/health/live').status_code==200
    assert client.get('/health/ready').json()['status']=='ready'
    assert 'JCF Enterprise' in client.get('/admin').text
    assert client.get('/assets/admin.js').status_code==200
    assert client.get('/api/openapi.json').status_code==200

def test_fail_closed_auth(client,app):
    from jcf_platform.main import _principal
    app.dependency_overrides.pop(_principal)
    assert client.get('/api/v1/registry/locations').status_code==401
    app.dependency_overrides[_principal]=lambda: {'sub':'wrong','_scopes':set()}
    assert client.get('/api/v1/registry/locations').status_code==403

def test_full_onboarding_flow(client,service_client):
    a=getid(service_client('LOCATION_CREATE',loc('DIV-A')))
    cib=getid(service_client('SECTION_DEFINE',section('CIB','CIB')))
    ls=getid(service_client('SECTION_ATTACH',attach(a,cib)))
    locations=client.get('/api/v1/registry/locations?asOf=2026-10-28').json()['items']
    assert locations[0]['locationId']==a
    assert not any(x.get('isDefault') for x in client.get(f'/api/v1/registry/locations/{a}/sections').json()['items'])
    assert client.get(f'/api/v1/registry/locations/{a}/sections').json()['items'][0]['name']=='CIB'
    v=client.post('/api/v1/registry/posting-selections/validate',json={'locationId':a,'sectionId':ls,'effectiveOn':'2026-10-28'})
    assert v.status_code==200
    b=getid(service_client('LOCATION_CREATE',loc('DIV-B','Division B')))
    bad=client.post('/api/v1/registry/posting-selections/validate',json={'locationId':b,'sectionId':ls,'effectiveOn':'2026-10-28'})
    assert bad.status_code==422 and bad.json()['code']=='SECTION_NOT_IN_LOCATION'
    other=getid(service_client('SECTION_ATTACH',attach(b,cib)))
    assert ls!=other
    assert client.get(f'/api/v1/registry/locations/{b}/sections').json()['items'][0]['sectionId']==other

def test_governance_rejects_self_approval(client,app):
    r=client.post('/api/v1/governance/requests',json={'action':'LOCATION_CREATE','payload':loc('A'),'reference':'ORDER-X'})
    rid=r.json()['changeRequestId']
    r=client.post(f'/api/v1/governance/requests/{rid}/decision',json={'decision':'APPROVED'})
    assert r.status_code==403 and r.json()['code']=='SELF_APPROVAL'
    r=client.post(f'/api/v1/governance/requests/{rid}/apply')
    assert r.status_code==409
    app.state.claims['sub']='another'
    assert client.post(f'/api/v1/governance/requests/{rid}/decision',json={'decision':'APPROVED'}).status_code==200
    assert client.post(f'/api/v1/governance/requests/{rid}/apply').status_code==200
    assert client.post(f'/api/v1/governance/requests/{rid}/apply').status_code==409

def test_name_history_and_alias(client,service_client):
    a=getid(service_client('LOCATION_CREATE',loc('SAC','Saint Andrew Central')))
    getid(service_client('LOCATION_ALIAS',dict(locationId=a,alias='SAC HQ')))
    assert len(client.get('/api/v1/registry/locations?search=SAC%20HQ').json()['items'])==1
    getid(service_client('LOCATION_RENAME',dict(locationId=a,name='New Central',effectiveFrom='2026-07-01')))
    old=client.get('/api/v1/registry/locations?asOf=2026-06-30').json()['items'][0]['name']
    new=client.get('/api/v1/registry/locations?asOf=2026-07-01').json()['items'][0]['name']
    assert old=='Saint Andrew Central' and new=='New Central'

def test_two_authorities_and_direction(client,service_client):
    div=getid(service_client('LOCATION_CREATE',loc('DIV-A')))
    cib_hq=getid(service_client('LOCATION_CREATE',loc('CIBHQ','CIB HQ')))
    strat=getid(service_client('LOCATION_CREATE',loc('STRAT','Strategic Operations')))
    cib=getid(service_client('SECTION_DEFINE',section('CIB','CIB')))
    unit=getid(service_client('SECTION_ATTACH',attach(div,cib)))
    for source,target,typ in [(unit,div,'LOCAL'),(unit,cib_hq,'FUNCTIONAL'),(div,strat,'OPERATIONAL')]:
        getid(service_client('REPORTING_EDGE',dict(sourceId=source,authorityId=target,relationType=typ,
                                                   validFrom='2020-01-01')))
    links=client.get(f'/api/v1/org-nodes/{unit}/reporting').json()['authorities']
    assert {(x['authorityOrgNodeId'],x['type']) for x in links}=={(div,'LOCAL'),(cib_hq,'FUNCTIONAL')}
    assert client.get(f'/api/v1/org-nodes/{cib_hq}/reporting').json()['authorities']==[]
    dup=service_client('REPORTING_EDGE',dict(sourceId=unit,authorityId=cib_hq,relationType='FUNCTIONAL',validFrom='2021-01-01'))
    assert dup.status_code==409

def test_temporal_parent_cycle(client,service_client):
    a=getid(service_client('LOCATION_CREATE',loc('A')))
    b=getid(service_client('LOCATION_CREATE',loc('B')))
    c=getid(service_client('LOCATION_CREATE',loc('C')))
    getid(service_client('LOCATION_PARENT',dict(childId=b,parentId=a,validFrom='2020-01-01')))
    getid(service_client('LOCATION_PARENT',dict(childId=c,parentId=b,validFrom='2020-01-01')))
    assert service_client('LOCATION_PARENT',dict(childId=a,parentId=c,validFrom='2020-01-01')).status_code==409
    assert service_client('LOCATION_PARENT',dict(childId=b,parentId=c,validFrom='2020-01-01')).status_code==409

def test_facilities_multi_building(client,service_client):
    locid=getid(service_client('LOCATION_CREATE',loc('A')))
    defin=getid(service_client('SECTION_DEFINE',section('CIB','CIB')))
    unit=getid(service_client('SECTION_ATTACH',attach(locid,defin)))
    lids=[]
    for name in 'ABC':
        b=getid(service_client('BUILDING_CREATE',dict(code='B'+name,name='Building '+name,address='Street '+name,
                        parish='Kingston',validFrom='2020-01-01',latitude=18.01,longitude=-76.81)))
        lids.append(getid(service_client('LOCATION_BUILDING',dict(locationId=locid,buildingId=b,validFrom='2020-01-01'))))
    for i,lid in enumerate(lids[:2]):
        getid(service_client('SECTION_BUILDING',dict(sectionId=unit,locationBuildingId=lid,
                                           validFrom='2020-01-01',assignmentType='PRIMARY' if i==0 else 'SECONDARY')))
    got=client.get(f'/api/v1/registry/sections/{unit}/buildings').json()['items']
    assert {b['name'] for b in got}=={'Building A','Building B'}
    assert len(client.get(f'/api/v1/registry/locations/{locid}/buildings').json()['items'])==3
    assert service_client('SECTION_BUILDING',dict(sectionId=unit,locationBuildingId=lids[1],
                         validFrom='2020-01-01',assignmentType='PRIMARY')).status_code==409

def test_gps_and_occupancy_validation(client,service_client):
    missing=service_client('BUILDING_CREATE',dict(code='BAD',name='Bad',address='X',parish='Y',
                              validFrom='2020-01-01',latitude=20.0))
    assert missing.status_code==422
    invalid=service_client('BUILDING_CREATE',dict(code='BAD2',name='Bad2',address='X',parish='Y',
                              validFrom='2020-01-01',latitude=92.0,longitude=-76.0))
    assert invalid.status_code==422

def test_personnel_posting_and_leave(client,service_client):
    locid=getid(service_client('LOCATION_CREATE',loc('A')))
    defin=getid(service_client('SECTION_DEFINE',section('CIB','CIB')))
    unit=getid(service_client('SECTION_ATTACH',attach(locid,defin)))
    p=getid(service_client('PERSONNEL_CREATE',dict(computerNumber='12345',displayName='Test Officer')))
    getid(service_client('POSTING_CREATE',dict(personnelId=p,locationId=locid,sectionId=unit,validFrom='2020-01-01')))
    assert service_client('POSTING_CREATE',dict(personnelId=p,locationId=locid,sectionId=unit,
                             validFrom='2022-01-01')).status_code==409
    assert len(client.get(f'/api/v1/hr/personnel/{p}/postings').json()['items'])==1
    getid(service_client('LEAVE_REQUEST',dict(personnelId=p,category='ANNUAL',startsOn='2026-11-01',endsBefore='2026-11-05')))
    assert client.get(f'/api/v1/hr/personnel/{p}/leave').json()['items'][0]['status']=='REQUESTED'

def test_jcfadreset_current_posting_contract(client,service_client,app):
    """A member's approved posting resolves to the EMP IDs, never a local duplicate."""
    a=getid(service_client('LOCATION_CREATE',loc('SAC','Saint Andrew Central')))
    b=getid(service_client('LOCATION_CREATE',loc('SAS','Saint Andrew South')))
    unit=getid(service_client('SECTION_DEFINE',section('CIB','CIB')))
    section_a=getid(service_client('SECTION_ATTACH',attach(a,unit)))
    section_b=getid(service_client('SECTION_ATTACH',attach(b,unit)))
    member=getid(service_client('PERSONNEL_CREATE',dict(computerNumber='77883',displayName='Example Member')))
    endpoint='/api/v1/hr/personnel/by-computer-number/77883/posting/current?asOf=2026-10-29'
    # A person with no approved primary posting must NOT be treated as assigned.
    empty=client.get(endpoint)
    assert empty.status_code==200 and empty.json()['posting'] is None
    getid(service_client('POSTING_CREATE',dict(personnelId=member,locationId=a,
                            sectionId=section_a,validFrom='2026-10-01')))
    posted=client.get(endpoint)
    assert posted.status_code==200,posted.text
    assert posted.json()['computerNumber']=='77883'
    assert posted.json()['posting']['locationId']==a
    assert posted.json()['posting']['sectionId']==section_a
    assert posted.json()['posting']['sectionId']!=section_b
    assert client.get('/api/v1/hr/personnel/by-computer-number/99999/posting/current').status_code==404
    app.state.claims['_scopes'].discard('hr.postings.read')
    assert client.get(endpoint).status_code==403


def test_restricted_visibility(client,service_client,app):
    locid=getid(service_client('LOCATION_CREATE',loc('SECRET','Protected HQ',visibility='RESTRICTED')))
    defin=getid(service_client('SECTION_DEFINE',section('CIB','CIB')))
    unit=getid(service_client('SECTION_ATTACH',attach(locid,defin,visibility='RESTRICTED')))
    assert client.get('/api/v1/registry/locations?search=Protected').json()['items']==[]
    assert client.get(f'/api/v1/registry/locations/{locid}/sections').status_code==404
    assert client.post('/api/v1/registry/posting-selections/validate',json={'locationId':locid,'sectionId':unit,'effectiveOn':'2026-10-28'}).status_code==404
    app.state.claims['_scopes'].add('registry.restricted.read')
    assert client.get(f'/api/v1/registry/locations/{locid}/sections').status_code==200

def test_contact_privacy(client,service_client,app):
    locid=getid(service_client('LOCATION_CREATE',loc('A')))
    getid(service_client('CONTACT_ASSIGN',dict(entityKind='LOCATION',entityId=locid,kind='PHONE',
                       value='8765550100',purpose='RECEPTION',visibility='INTERNAL',validFrom='2020-01-01')))
    getid(service_client('CONTACT_ASSIGN',dict(entityKind='LOCATION',entityId=locid,kind='MOBILE',
                       value='8765559999',purpose='PRIVATE',visibility='PERSONAL',validFrom='2020-01-01')))
    assert len(client.get(f'/api/v1/contacts?entityKind=LOCATION&entityId={locid}').json()['items'])==1
    app.state.claims['_scopes'].add('contacts.personal.read')
    assert len(client.get(f'/api/v1/contacts?entityKind=LOCATION&entityId={locid}').json()['items'])==2

def test_batch_import_excel_preview(client):
    content=b'type,code,name,validFrom,locationType\nLOCATION,SAC,Saint Andrew Central,2020-01-01,DIVISION\nSECTION,CIB,CIB,2020-01-01,\n'
    r=client.post('/api/v1/import/preview',files={'file':('data.csv',content,'text/csv')})
    assert r.status_code==200,r.text
    assert len(r.json()['operations'])==2

def test_batch_import_transaction(client,service_client):
    ops=[{'action':'LOCATION_CREATE','payload':loc('DIV-A')},{'action':'SECTION_DEFINE','payload':section('CIB','CIB')}]
    getid(service_client('BATCH_REGISTER',dict(operations=ops)))
    assert len(client.get('/api/v1/registry/locations').json()['items'])==1
    bad=service_client('BATCH_REGISTER',dict(operations=[{'action':'LOCATION_CREATE','payload':loc('DIV-B')},
         {'action':'LOCATION_CREATE','payload':loc('DIV-A')}]))
    assert bad.status_code==409
    assert len(client.get('/api/v1/registry/locations').json()['items'])==1

def test_event_feed(client,service_client):
    getid(service_client('LOCATION_CREATE',loc('A')))
    events=client.get('/api/v1/registry/changes?after=0').json()['items']
    assert len(events)==1 and events[0]['kind']=='LOCATION_CREATE'

def test_date_overlap_algorithm():
    from jcf_platform.service import overlap,contained
    assert not overlap(date(2020,1,1),date(2020,2,1),date(2020,2,1),None)
    assert overlap(date(2020,1,1),None,date(2026,1,1),None)
    assert not contained(date(2020,1,1),None,date(2020,1,1),date(2021,1,1))

def test_batch_cannot_escalate_to_hr(client,service_client):
    body={'operations':[{'action':'PERSONNEL_CREATE','payload':{'computerNumber':'12345','displayName':'Intruder'}}]}
    r=service_client('BATCH_REGISTER',body)
    assert r.status_code==403 and r.json()['code']=='INVALID_BATCH'

def test_section_restricted_filter_and_reporting_target(client,service_client,app):
    locid=getid(service_client('LOCATION_CREATE',loc('A')))
    secret=getid(service_client('LOCATION_CREATE',loc('B','Secret HQ',visibility='RESTRICTED')))
    d=getid(service_client('SECTION_DEFINE',section('CIB','CIB')))
    restricted=getid(service_client('SECTION_ATTACH',attach(locid,d,visibility='RESTRICTED')))
    getid(service_client('REPORTING_EDGE',dict(sourceId=locid,authorityId=secret,
                                                relationType='FUNCTIONAL',validFrom='2020-01-01')))
    assert client.get(f'/api/v1/registry/locations/{locid}/sections').json()['items']==[]
    assert client.get(f'/api/v1/org-nodes/{locid}/reporting').json()['authorities']==[]
    app.state.claims['_scopes'].add('registry.restricted.read')
    assert client.get(f'/api/v1/registry/locations/{locid}/sections').json()['items'][0]['sectionId']==restricted
    assert len(client.get(f'/api/v1/org-nodes/{locid}/reporting').json()['authorities'])==1

def test_schema_invalid_request_error(client):
    r=client.post('/api/v1/registry/posting-selections/validate',json={'locationId':'invalid'})
    assert r.status_code==422 and r.headers['content-type']=='application/problem+json'
    assert r.json()['code']=='INVALID_REQUEST'

def test_restricted_contact_owner_is_not_disclosed(client,service_client):
    hidden=getid(service_client('LOCATION_CREATE',loc('SH','Secret HQ',visibility='RESTRICTED')))
    getid(service_client('CONTACT_ASSIGN',dict(entityKind='LOCATION',entityId=hidden,kind='PHONE',value='0123',
                      purpose='RECEPTION',validFrom='2020-01-01')))
    r=client.get(f'/api/v1/contacts?entityKind=LOCATION&entityId={hidden}')
    assert r.status_code==404

def test_scope_by_domain(client,service_client,app):
    app.state.claims['_scopes'].discard('hr.write')
    r=client.post('/api/v1/governance/requests',json={'action':'PERSONNEL_CREATE',
         'payload':{'computerNumber':'12345','displayName':'Officer'},'reference':'ORDER-A'})
    assert r.status_code==403
    r=client.post('/api/v1/governance/requests',json={'action':'LOCATION_CREATE','payload':loc('A'),
                                                    'reference':'ORDER-A'})
    assert r.status_code==201

def test_request_idempotency_prevents_duplicate(client):
    key='replay-unique-test-12345'
    body={'action':'LOCATION_CREATE','payload':loc('A'),'reference':'ORDER-A'}
    first=client.post('/api/v1/governance/requests',json=body,headers={'Idempotency-Key':key})
    again=client.post('/api/v1/governance/requests',json=body,headers={'Idempotency-Key':key})
    assert first.status_code==201 and first.json()['changeRequestId']==again.json()['changeRequestId']
    bad=client.post('/api/v1/governance/requests',json={**body,'reference':'ORDER-B'},
                    headers={'Idempotency-Key':key})
    assert bad.status_code==409 and bad.json()['code']=='IDEMPOTENCY_CONFLICT'

def test_sqlalchemy_schema_has_fk_guards():
    from jcf_platform.models import Base
    from sqlalchemy import ForeignKeyConstraint
    sec=Base.metadata.tables['section_building']
    assert len([c for c in sec.constraints if isinstance(c,ForeignKeyConstraint)])==2
    assert 'idempotency_key' in Base.metadata.tables['change_request'].columns

def test_import_xlsx(client):
    from io import BytesIO
    from openpyxl import Workbook
    book=Workbook();ws=book.active
    ws.append(['type','code','name','validFrom','locationType'])
    ws.append(['LOCATION','SAC','Saint Andrew Central','2020-01-01','DIVISION'])
    output=BytesIO();book.save(output)
    r=client.post('/api/v1/import/preview',files={'file':('locations.xlsx',output.getvalue(),
                                                       'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert r.status_code==200,r.text
    assert r.json()['operations'][0]['payload']['code']=='SAC'

def test_excel_code_based_section_mapping(client,service_client):
    csv_bytes=(b'type,code,name,validFrom,locationType,locationCode,sectionCode\n'
             b'MAPPING,,,2020-01-01,,DIV-A,CIB\n'
             b'SECTION,CIB,CIB,2020-01-01,,,\n'
             b'LOCATION,DIV-A,Division A,2020-01-01,DIVISION,,\n')
    preview=client.post('/api/v1/import/preview',files={'file':('data.csv',csv_bytes,'text/csv')})
    assert preview.status_code==200,preview.text
    assert len(preview.json()['operations'])==3 and not preview.json()['errors']
    r=service_client('BATCH_REGISTER',{'operations':preview.json()['operations']})
    assert r.status_code==200,r.text
    loc_id=client.get('/api/v1/registry/locations').json()['items'][0]['locationId']
    assert client.get(f'/api/v1/registry/locations/{loc_id}/sections').json()['items'][0]['name']=='CIB'

def test_section_rename_keeps_instance_ids(client,service_client):
    locid=getid(service_client('LOCATION_CREATE',loc('DIV-A')))
    defin=getid(service_client('SECTION_DEFINE',section('CIB','CIB')))
    unit=getid(service_client('SECTION_ATTACH',attach(locid,defin)))
    getid(service_client('SECTION_RENAME',dict(sectionDefinitionId=defin,name='Criminal Investigations Branch',effectiveFrom='2026-01-01')))
    before=client.get(f'/api/v1/registry/locations/{locid}/sections?asOf=2025-01-01').json()['items'][0]
    after=client.get(f'/api/v1/registry/locations/{locid}/sections?asOf=2026-02-01').json()['items'][0]
    assert before['name']=='CIB' and after['name']=='Criminal Investigations Branch'
    assert before['sectionId']==after['sectionId']==unit

def test_restricted_retirement_guard_and_asof(client,service_client):
    locid=getid(service_client('LOCATION_CREATE',loc('A')))
    defin=getid(service_client('SECTION_DEFINE',section('CIB','CIB')))
    unit=getid(service_client('SECTION_ATTACH',attach(locid,defin)))
    assert service_client('LOCATION_RETIRE',dict(locationId=locid,effectiveFrom='2026-01-01')).status_code==409
    getid(service_client('SECTION_RETIRE',dict(sectionId=unit,effectiveFrom='2026-01-01')))
    getid(service_client('LOCATION_RETIRE',dict(locationId=locid,effectiveFrom='2026-01-01')))
    assert client.get('/api/v1/registry/locations?asOf=2025-12-31').json()['items'][0]['locationId']==locid
    assert client.get('/api/v1/registry/locations?asOf=2026-01-01').json()['items']==[]

def test_alias_search_for_section(client,service_client):
    locid=getid(service_client('LOCATION_CREATE',loc('A')))
    defin=getid(service_client('SECTION_DEFINE',section('CIB','Criminal Investigations Branch')))
    getid(service_client('SECTION_ATTACH',attach(locid,defin)))
    getid(service_client('SECTION_ALIAS',dict(definitionId=defin,alias='CIB')))
    assert len(client.get(f'/api/v1/registry/locations/{locid}/sections?search=cib').json()['items'])==1
    assert client.get(f'/api/v1/registry/locations/{locid}/sections?search=notfound').json()['items']==[]

def test_openapi_contract_declares_bearer(client):
    contract=client.get('/api/openapi.json').json()
    assert 'OIDC_Bearer' in contract['components']['securitySchemes']
    assert contract['paths']['/api/v1/registry/locations']['get']['security']

def test_administration_search_sources(client,service_client):
    l=getid(service_client('LOCATION_CREATE',loc('DIV-A')))
    d=getid(service_client('SECTION_DEFINE',section('CIB','CIB')))
    ls=getid(service_client('SECTION_ATTACH',attach(l,d)))
    defs=client.get('/api/v1/registry/section-definitions?search=CIB').json()['items']
    assert defs[0]['sectionDefinitionId']==d
    nodes=client.get('/api/v1/registry/org-nodes?search=Division').json()['items']
    assert {x['orgNodeId'] for x in nodes}=={l,ls}

def test_facility_lookup_sources(client,service_client):
    l=getid(service_client('LOCATION_CREATE',loc('A')))
    b=getid(service_client('BUILDING_CREATE',dict(code='BUILD-A',name='North Block',
           address='Street 1',parish='Kingston',validFrom='2020-01-01')))
    lb=getid(service_client('LOCATION_BUILDING',dict(locationId=l,buildingId=b,validFrom='2020-01-01')))
    assert client.get('/api/v1/facilities/buildings?search=North').json()['items'][0]['buildingId']==b
    assert client.get(f'/api/v1/facilities/location-building-links?locationId={l}').json()['items'][0]['locationBuildingId']==lb

def test_oidc_signature_expiry_and_audience(client,app,monkeypatch):
    from datetime import datetime, timezone, timedelta
    from cryptography.hazmat.primitives.asymmetric import rsa
    import jwt
    import jcf_platform.main as main
    from jcf_platform.main import _principal
    app.dependency_overrides.pop(_principal,None)
    app.state.settings.update({'issuer':'https://issuer.test','audience':'jcf-registry','jwks':'https://issuer.test/keys'})
    key=rsa.generate_private_key(public_exponent=65537,key_size=2048)
    class Key:
        def __init__(self):self.key=key.public_key()
    class Jwks:
        def get_signing_key_from_jwt(self,token):return Key()
    monkeypatch.setattr(main,'jwk_client',lambda _url:Jwks())
    now=datetime.now(timezone.utc)
    claims={'sub':'test-client','iss':'https://issuer.test','aud':'jcf-registry',
            'scope':'registry.read','exp':now+timedelta(minutes=2)}
    valid=jwt.encode(claims,key,algorithm='RS256')
    assert client.get('/api/v1/registry/locations',headers={'Authorization':'Bearer '+valid}).status_code==200
    expired=jwt.encode({**claims,'exp':now-timedelta(hours=1)},key,algorithm='RS256')
    assert client.get('/api/v1/registry/locations',headers={'Authorization':'Bearer '+expired}).status_code==401
    audience=jwt.encode({**claims,'aud':'other-service'},key,algorithm='RS256')
    assert client.get('/api/v1/registry/locations',headers={'Authorization':'Bearer '+audience}).status_code==401
    other_key=rsa.generate_private_key(public_exponent=65537,key_size=2048)
    forged=jwt.encode(claims,other_key,algorithm='RS256')
    assert client.get('/api/v1/registry/locations',headers={'Authorization':'Bearer '+forged}).status_code==401

def test_sqlite_restart_persists_registry(tmp_path):
    from fastapi.testclient import TestClient
    from jcf_platform.main import create_app,_principal
    from jcf_platform.models import Base
    file=tmp_path/'jcf_registry_test.sqlite'
    app=create_app('sqlite:///'+str(file));Base.metadata.create_all(app.state.engine)
    app.state.claims={'sub':'maker','_scopes':{'registry.read','registry.write','registry.approve','registry.apply'}}
    app.dependency_overrides[_principal]=lambda:app.state.claims
    with TestClient(app) as c:
        req=c.post('/api/v1/governance/requests',json={'action':'LOCATION_CREATE','payload':loc('A'),
                              'reference':'TEST-RESTORE'})
        crid=req.json()['changeRequestId']
        app.state.claims['sub']='approver'
        assert c.post(f'/api/v1/governance/requests/{crid}/decision',json={'decision':'APPROVED'}).status_code==200
        assert c.post(f'/api/v1/governance/requests/{crid}/apply').status_code==200
    restored=create_app('sqlite:///'+str(file))
    restored.dependency_overrides[_principal]=lambda:{'sub':'reader','_scopes':{'registry.read'}}
    with TestClient(restored) as c:
        assert c.get('/api/v1/registry/locations').json()['items'][0]['code']=='A'
